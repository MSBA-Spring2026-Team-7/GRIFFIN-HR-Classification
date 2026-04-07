"""Build griffin-task-tracker.xlsx -- team assignment tracker."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from datetime import date, timedelta

wb = Workbook()

# Colors
HDR_FILL = PatternFill("solid", fgColor="343a40")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
DONE_FILL = PatternFill("solid", fgColor="d4edda")
IP_FILL = PatternFill("solid", fgColor="fff3cd")
BLOCKED_FILL = PatternFill("solid", fgColor="f8d7da")
SECTION_FILL = PatternFill("solid", fgColor="e9ecef")
DATA_FONT = Font(name="Arial", size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="dee2e6"),
    right=Side(style="thin", color="dee2e6"),
    top=Side(style="thin", color="dee2e6"),
    bottom=Side(style="thin", color="dee2e6"),
)

tasks = [
    (1, "Project spec document (v1-v3)", "Done", "", "Medium", "Shared", "WS1", "2026-02-24", "2026-03-07", "Foundation document for both courses", ""),
    (2, "Database schema design (9 tables)", "Done", "", "Medium", "Shared", "WS1", "2026-03-01", "2026-03-07", "6 reference + 1 crosswalk + 2 output tables", "Project spec"),
    (3, "Proof of Concept prototype", "Done", "", "Medium", "AI Course", "WS1", "2026-03-03", "2026-03-09", "Early prototype validating approach", ""),
    (4, "Scrape 56 DHRM pages", "Done", "", "High", "Big Data", "WS1", "2026-03-10", "2026-03-15", "50 HTML + 6 PDF cached in raw_pages/", ""),
    (5, "Parse HTML (50 groups)", "Done", "", "High", "Big Data", "WS1", "2026-03-12", "2026-03-18", "BeautifulSoup + regex extraction", "Scrape"),
    (6, "Parse PDF (6 groups, 3 variants)", "Done", "", "High", "Big Data", "WS1", "2026-03-16", "2026-03-24", "pdfplumber, 3 format variants", "Scrape"),
    (7, "Classification prompt v1", "Done", "", "Medium", "Shared", "WS2", "2026-03-20", "2026-03-24", "Draft with blend logic + explanation", "Reference data"),
    (8, "Data validation (56 groups, 294 roles)", "Done", "", "High", "Big Data", "WS1", "2026-03-22", "2026-03-24", "100% coverage verified", "Parse HTML + PDF"),
    (9, "FY26 salary + W&M grades + crosswalk", "Done", "", "High", "Shared", "WS1", "2026-03-23", "2026-03-25", "9 bands + 46 grades + Band 7->S18", ""),
    (10, "DB load + Excel export", "Done", "", "High", "Big Data", "WS1", "2026-03-25", "2026-03-26", "7 tables SQLite + 7 .xlsx in hr_handoff/", "Parse + Salary"),
    (11, "URL refactoring in notebook", "In Progress", "Salva", "Medium", "Big Data", "WS1", "2026-03-27", "2026-03-28", "Base URL + f-strings", ""),
    (12, "Notebook gap analysis vs. rubrics", "In Progress", "Salva", "Medium", "Shared", "WS1", "2026-03-27", "2026-03-29", "Review against both course rubrics", ""),
    (13, "Refine classification prompt v2", "In Progress", "Salva", "Medium", "Shared", "WS2", "2026-03-28", "2026-04-01", "Test against Assistant Controller", "Prompt v1"),
    (14, "Deploy to Google Cloud SQL", "To Do", "", "High", "Big Data", "WS3", "2026-03-28", "2026-04-01", "Migrate SQLite to MySQL on GCP", "DB load"),
    (15, "GitHub repo + README draft", "To Do", "", "High", "AI Course", "WS4", "2026-03-28", "2026-04-02", "7% rubric: repo organization", ""),
    (16, "GitHub Projects board setup", "To Do", "", "High", "AI Course", "WS4", "2026-03-28", "2026-03-29", "7% rubric: kanban board", "GitHub repo"),
    (17, "Research paper search", "To Do", "", "Medium", "AI Course", "WS4", "2026-03-29", "2026-04-02", "7% rubric: paper quality + discussion", ""),
    (18, "Streamlit app skeleton", "To Do", "", "High", "AI Course", "WS3", "2026-03-29", "2026-04-02", "Basic structure, no LLM yet", "Excel files"),
    (19, "Streamlit: LLM integration", "To Do", "", "High", "AI Course", "WS3", "2026-04-03", "2026-04-07", "Connect to LLM API, agentic design", "Skeleton + Prompt v2"),
    (20, "Streamlit: deploy to URL", "To Do", "", "High", "AI Course", "WS3", "2026-04-07", "2026-04-09", "Public URL for live demo", "LLM integration"),
    (21, "Notebook visualizations", "To Do", "", "High", "Big Data", "WS1", "2026-04-03", "2026-04-06", "matplotlib/seaborn/plotly, 4+ charts", "DataFrames"),
    (22, "Pay estimation prompt", "To Do", "", "Medium", "Shared", "WS2", "2026-04-02", "2026-04-04", "Weighted pay math for blends", "Prompt v2 + Crosswalk"),
    (23, "Validate 10+ W&M positions", "To Do", "", "High", "Shared", "WS4", "2026-04-03", "2026-04-09", "Real postings from Workday", "Prompt v2 + Pay prompt"),
    (24, "Responsible AI write-up", "To Do", "", "Medium", "AI Course", "WS4", "2026-04-05", "2026-04-07", "README + poster section", ""),
    (25, "Ethical reflection", "To Do", "", "Medium", "Big Data", "WS4", "2026-04-05", "2026-04-07", "Data ethics for notebook + slides", ""),
    (26, "Project folder structure", "To Do", "", "Medium", "Big Data", "WS4", "2026-04-06", "2026-04-07", "data/ notebooks/ reports/ ppt/", ""),
    (27, "Poster design + content", "To Do", "", "High", "AI Course", "WS4", "2026-04-08", "2026-04-12", "7% rubric: poster + 6% presentation", "Streamlit + Research + RAI"),
    (28, "README finalize", "To Do", "", "High", "AI Course", "WS4", "2026-04-10", "2026-04-13", "Screenshots, refs, final polish", "Streamlit deployed"),
    (29, "Dry run: AI poster presentation", "To Do", "", "High", "AI Course", "WS4", "2026-04-13", "2026-04-14", "Full team rehearsal", "Poster + Streamlit"),
    (30, "Big Data slide deck", "To Do", "", "High", "Big Data", "WS4", "2026-04-10", "2026-04-15", "Problem, architecture, demo, findings", "Viz + Cloud SQL + Ethics"),
    (31, "Dry run: Big Data presentation", "To Do", "", "High", "Big Data", "WS4", "2026-04-15", "2026-04-16", "Full team rehearsal", "Slide deck"),
    (32, "Data dictionary", "To Do", "", "Medium", "HR Handoff", "WS4", "2026-04-11", "2026-04-14", "Every field documented for HR", "Excel files"),
    (33, "User guide / SOP", "To Do", "", "Medium", "HR Handoff", "WS4", "2026-04-12", "2026-04-16", "Step-by-step for non-technical HR", "Prompts + Streamlit"),
    (34, "Validation report", "To Do", "", "Medium", "HR Handoff", "WS4", "2026-04-14", "2026-04-17", "Test results + accuracy metrics", "Validate 10+"),
]

# ===== SHEET 1: Task Board =====
ws = wb.active
ws.title = "Task Board"
ws.sheet_properties.tabColor = "343a40"

headers = ["#", "Task", "Status", "Owner", "Priority", "Course", "Workstream", "Start", "End", "Notes / Rubric", "Dependencies"]
col_widths = [5, 38, 14, 16, 10, 14, 13, 12, 12, 42, 28]

for c, (hdr, w) in enumerate(zip(headers, col_widths), 1):
    cell = ws.cell(row=1, column=c, value=hdr)
    cell.font = HDR_FONT
    cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER
    ws.column_dimensions[get_column_letter(c)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{len(tasks)+1}"

for r, t in enumerate(tasks, 2):
    for c, val in enumerate(t, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        if c in (1, 3, 5, 6, 7, 8, 9):
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.alignment = Alignment(vertical="center", wrap_text=True)

# Data validation dropdowns
dv_status = DataValidation(type="list", formula1='"Done,In Progress,To Do,Blocked"', allow_blank=True)
ws.add_data_validation(dv_status)
dv_status.add(f"C2:C{len(tasks)+1}")

dv_priority = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
ws.add_data_validation(dv_priority)
dv_priority.add(f"E2:E{len(tasks)+1}")

dv_course = DataValidation(type="list", formula1='"AI Course,Big Data,HR Handoff,Shared"', allow_blank=True)
ws.add_data_validation(dv_course)
dv_course.add(f"F2:F{len(tasks)+1}")

dv_ws = DataValidation(type="list", formula1='"WS1,WS2,WS3,WS4"', allow_blank=True)
ws.add_data_validation(dv_ws)
dv_ws.add(f"G2:G{len(tasks)+1}")

# Conditional formatting -- entire row colors by status
last_row = len(tasks) + 1
ws.conditional_formatting.add(f"A2:K{last_row}",
    FormulaRule(formula=['$C2="Done"'], fill=DONE_FILL))
ws.conditional_formatting.add(f"A2:K{last_row}",
    FormulaRule(formula=['$C2="In Progress"'], fill=IP_FILL))
ws.conditional_formatting.add(f"A2:K{last_row}",
    FormulaRule(formula=['$C2="Blocked"'], fill=BLOCKED_FILL))

# ===== SHEET 2: Team Dashboard =====
ds = wb.create_sheet("Team Dashboard")
ds.sheet_properties.tabColor = "17a2b8"

TITLE_FONT = Font(name="Arial", bold=True, size=14, color="343a40")
SUBTITLE_FONT = Font(name="Arial", bold=True, size=11, color="495057")
METRIC_FONT = Font(name="Arial", bold=True, size=20, color="343a40")
LABEL_FONT = Font(name="Arial", size=10, color="6c757d")

ds["A1"] = "GRIFFIN Project Dashboard"
ds["A1"].font = TITLE_FONT
ds.merge_cells("A1:F1")

ds["A3"] = "Status Summary"
ds["A3"].font = SUBTITLE_FONT
status_items = [
    ("Done", '=COUNTIF(\'Task Board\'!C:C,"Done")', DONE_FILL),
    ("In Progress", '=COUNTIF(\'Task Board\'!C:C,"In Progress")', IP_FILL),
    ("To Do", '=COUNTIF(\'Task Board\'!C:C,"To Do")', PatternFill("solid", fgColor="e3f2fd")),
    ("Blocked", '=COUNTIF(\'Task Board\'!C:C,"Blocked")', BLOCKED_FILL),
    ("Total", "=COUNTA('Task Board'!C2:C100)", SECTION_FILL),
]
for c, (label, formula, fill) in enumerate(status_items, 1):
    ds.cell(row=4, column=c, value=label).font = LABEL_FONT
    ds.cell(row=4, column=c).alignment = Alignment(horizontal="center")
    val_cell = ds.cell(row=5, column=c, value=formula)
    val_cell.font = METRIC_FONT
    val_cell.alignment = Alignment(horizontal="center")
    val_cell.fill = fill
    ds.column_dimensions[get_column_letter(c)].width = 16

# Course breakdown
ds["A7"] = "Tasks by Course"
ds["A7"].font = SUBTITLE_FONT
for r, (label, formula) in enumerate([
    ("AI Course", '=COUNTIF(\'Task Board\'!F:F,"AI Course")'),
    ("Big Data", '=COUNTIF(\'Task Board\'!F:F,"Big Data")'),
    ("HR Handoff", '=COUNTIF(\'Task Board\'!F:F,"HR Handoff")'),
    ("Shared", '=COUNTIF(\'Task Board\'!F:F,"Shared")'),
], 8):
    ds.cell(row=r, column=1, value=label).font = DATA_FONT
    ds.cell(row=r, column=2, value=formula).font = Font(name="Arial", bold=True, size=11)

# Owner workload
ds["A13"] = "Owner Workload (type names in yellow cells)"
ds["A13"].font = SUBTITLE_FONT
owner_hdrs = ["Owner Name", "Total", "To Do", "In Progress", "Done"]
for c, h in enumerate(owner_hdrs, 1):
    cell = ds.cell(row=14, column=c, value=h)
    cell.font = Font(name="Arial", bold=True, size=10)
    cell.border = THIN_BORDER

for r in range(15, 21):
    ds.cell(row=r, column=1).fill = PatternFill("solid", fgColor="fff3cd")
    ds.cell(row=r, column=1).border = THIN_BORDER
    ds.cell(row=r, column=2, value=f'=IF(A{r}="","",COUNTIF(\'Task Board\'!D:D,A{r}))').border = THIN_BORDER
    ds.cell(row=r, column=3, value=f'=IF(A{r}="","",COUNTIFS(\'Task Board\'!D:D,A{r},\'Task Board\'!C:C,"To Do"))').border = THIN_BORDER
    ds.cell(row=r, column=4, value=f'=IF(A{r}="","",COUNTIFS(\'Task Board\'!D:D,A{r},\'Task Board\'!C:C,"In Progress"))').border = THIN_BORDER
    ds.cell(row=r, column=5, value=f'=IF(A{r}="","",COUNTIFS(\'Task Board\'!D:D,A{r},\'Task Board\'!C:C,"Done"))').border = THIN_BORDER

# Quick lookup
ds["A23"] = "Quick Lookup: My Tasks"
ds["A23"].font = SUBTITLE_FONT
ds["A24"] = "Type your name:"
ds["A24"].font = DATA_FONT
ds["B24"].fill = PatternFill("solid", fgColor="fff3cd")
ds["A25"] = "Your open tasks:"
ds["A25"].font = DATA_FONT
ds["B25"] = '=IF(B24="","",COUNTIFS(\'Task Board\'!D:D,B24,\'Task Board\'!C:C,"<>Done"))'
ds["B25"].font = METRIC_FONT

# Key dates
ds["D7"] = "Key Dates"
ds["D7"].font = SUBTITLE_FONT
for r, line in enumerate(["AI Poster: April 15, 2026", "Big Data Deck: April 17, 2026", "Client status: Awaiting response"], 8):
    ds.cell(row=r, column=4, value=line).font = DATA_FONT

# ===== SHEET 3: Timeline =====
ts = wb.create_sheet("Timeline")
ts.sheet_properties.tabColor = "8e44ad"

week_starts = []
d = date(2026, 2, 23)
while d <= date(2026, 4, 20):
    week_starts.append(d)
    d += timedelta(days=7)

ts.cell(row=1, column=1, value="Task").font = HDR_FONT
ts.cell(row=1, column=1).fill = HDR_FILL
ts.cell(row=1, column=1).border = THIN_BORDER
ts.column_dimensions["A"].width = 38

ts.cell(row=1, column=2, value="Status").font = HDR_FONT
ts.cell(row=1, column=2).fill = HDR_FILL
ts.cell(row=1, column=2).border = THIN_BORDER
ts.column_dimensions["B"].width = 13

for c, ws_date in enumerate(week_starts, 3):
    cell = ts.cell(row=1, column=c, value=ws_date.strftime("%b %d"))
    cell.font = HDR_FONT
    cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center")
    cell.border = THIN_BORDER
    ts.column_dimensions[get_column_letter(c)].width = 9

BAR_DONE = PatternFill("solid", fgColor="28a745")
BAR_IP = PatternFill("solid", fgColor="ffc107")
BAR_AI = PatternFill("solid", fgColor="17a2b8")
BAR_BD = PatternFill("solid", fgColor="8e44ad")
BAR_HR = PatternFill("solid", fgColor="27ae60")
BAR_SHARED = PatternFill("solid", fgColor="e67e22")

ts.freeze_panes = "C2"

for r, t in enumerate(tasks, 2):
    num, name, status, owner, pri, course, wsm, start_s, end_s, notes, deps = t
    ts.cell(row=r, column=1, value=name).font = DATA_FONT
    ts.cell(row=r, column=1).border = THIN_BORDER
    ts.cell(row=r, column=2, value=status).font = DATA_FONT
    ts.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    ts.cell(row=r, column=2).border = THIN_BORDER

    start_d = date.fromisoformat(start_s)
    end_d = date.fromisoformat(end_s)

    for c, ws_date in enumerate(week_starts, 3):
        ws_end = ws_date + timedelta(days=6)
        cell = ts.cell(row=r, column=c)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")
        if start_d <= ws_end and end_d >= ws_date:
            cell.value = "\u2713" if status == "Done" else "X"
            cell.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
            if status == "Done":
                cell.fill = BAR_DONE
            elif status == "In Progress":
                cell.fill = BAR_IP
            else:
                fills_map = {"AI Course": BAR_AI, "Big Data": BAR_BD, "HR Handoff": BAR_HR, "Shared": BAR_SHARED}
                cell.fill = fills_map.get(course, BAR_SHARED)

# Milestone row
ms_row = len(tasks) + 3
ts.cell(row=ms_row, column=1, value="DEADLINES").font = Font(name="Arial", bold=True, size=10, color="dc3545")
for c, ws_date in enumerate(week_starts, 3):
    ws_end = ws_date + timedelta(days=6)
    labels = []
    if ws_date <= date(2026, 4, 15) <= ws_end:
        labels.append("AI Apr 15")
    if ws_date <= date(2026, 4, 17) <= ws_end:
        labels.append("BD Apr 17")
    if labels:
        cell = ts.cell(row=ms_row, column=c, value=" | ".join(labels))
        cell.font = Font(name="Arial", bold=True, size=8, color="dc3545")
        cell.fill = PatternFill("solid", fgColor="f8d7da")
        cell.alignment = Alignment(horizontal="center")

# Legend
lg = ms_row + 2
ts.cell(row=lg, column=1, value="Legend:").font = Font(name="Arial", bold=True, size=10)
for i, (label, fill) in enumerate([
    ("Done", BAR_DONE), ("In Progress", BAR_IP),
    ("AI Course", BAR_AI), ("Big Data", BAR_BD),
    ("HR Handoff", BAR_HR), ("Shared", BAR_SHARED),
]):
    ts.cell(row=lg + 1 + i, column=1, value=label).font = DATA_FONT
    ts.cell(row=lg + 1 + i, column=2).fill = fill

out = r"C:\Users\salva\Desktop\MSBA\Claude Skills\Courses\AI Course\HR_Classification_Project\docs\griffin-task-tracker.xlsx"
wb.save(out)
print(f"Saved: {out}")
